from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from uuid import uuid4

import openpyxl

try:
    from src.editor_v2 import compile_document_to_txt
except ModuleNotFoundError:
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from src.editor_v2 import compile_document_to_txt


RAW_ROOT = Path("data/raw")
ANNOTATION_SHEET_NAME = "\u6807\u6ce8\u8868"
NON_DIALOGUE_SPEECH_MODE = "non_dialogue"
TE_PATTERN = re.compile(r"<te\s+href=([\"']?)([^\"'>\s]+)\1>(.*?)</te>", re.IGNORECASE | re.DOTALL)
TAG_PATTERN = re.compile(r"</?[^>]+>")


TYPE_MAP = {
    "\u65f6\u95f4&\u5730\u70b9&\u4eba\u7269": "scene_cast",
    "\u65c1\u767d": "narration",
    "\u53f0\u8bcd": "line",
    "\u6587\u672c\u6807\u9898": "text_title",
    "\u6587\u672c\u6b63\u6587": "text_body",
    "\u5206\u652f\u9009\u9879": "branch_option",
    "\u5206\u652f\u540e\u7eed": "branch_followup",
    "\u672c\u7ae0\u6897\u6982": "chapter_summary",
    "\u6bb5\u843d\u6897\u6982": "paragraph_summary",
    "\u672c\u7ae0\u6ce8\u91ca": "chapter_note",
}


@dataclass
class Row:
    row_number: int
    source_id: str
    order: int
    raw_type: str
    block_type: str
    speaker: str
    raw_text: str
    keywords: list[str]
    section_index: int
    inline_note: str = ""


def new_id(prefix: str = "blk") -> str:
    return f"{prefix}_{uuid4().hex[:8]}"


def empty_annotations() -> dict[str, Any]:
    return {"manual_refs": [], "omissions": [], "settings": []}


def normalize_omissions(values: Any, text: str) -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []
    seen: set[tuple[str, int, int]] = set()
    text_value = str(text or "")
    for item in values or []:
        if not isinstance(item, dict):
            continue
        term = str(item.get("term", "")).strip()
        try:
            start = int(item.get("start"))
            end = int(item.get("end"))
        except (TypeError, ValueError):
            continue
        if not term or start < 0 or end <= start or end > len(text_value):
            continue
        # Only keep position-level cancellations that still point at the same text.
        if text_value[start:end] != term:
            continue
        key = (term, start, end)
        if key in seen:
            continue
        seen.add(key)
        normalized.append({"term": term, "start": start, "end": end})
    return normalized


def transferable_annotations(annotations: Any, text: str) -> dict[str, Any]:
    data = annotations if isinstance(annotations, dict) else {}
    return {
        "manual_refs": [],
        "omissions": normalize_omissions(data.get("omissions", []), text),
        "settings": [],
    }


def node_text_for_annotations(node: dict[str, Any]) -> str:
    if "label" in node:
        return str(node.get("label", ""))
    return str(node.get("text", ""))


def node_fingerprint(node: dict[str, Any]) -> tuple[str, str, str, str]:
    node_type = str(node.get("type", "option")).strip()
    speaker = str(node.get("speaker", "")).strip()
    text = node_text_for_annotations(node)
    source_id = str((node.get("meta") or {}).get("source_id", "")).strip()
    return (source_id, node_type, speaker, text)


def node_content_fingerprint(node: dict[str, Any]) -> tuple[str, str, str]:
    node_type = str(node.get("type", "option")).strip()
    speaker = str(node.get("speaker", "")).strip()
    text = node_text_for_annotations(node)
    return (node_type, speaker, text)


def build_annotation_index(nodes: list[dict[str, Any]]) -> dict[str, dict[tuple[Any, ...], list[dict[str, Any]]]]:
    index: dict[str, dict[tuple[Any, ...], list[dict[str, Any]]]] = {"source": {}, "content": {}}

    def add_node(node: dict[str, Any]) -> None:
        if not isinstance(node, dict):
            return
        text = node_text_for_annotations(node)
        annotations = transferable_annotations(node.get("annotations"), text)
        if annotations["omissions"]:
            source_key = node_fingerprint(node)
            content_key = node_content_fingerprint(node)
            if source_key[0]:
                index["source"].setdefault(source_key, []).append(annotations)
            # Fallback for editor files created before xlsx source IDs existed.
            index["content"].setdefault(content_key, []).append(annotations)
        for option in node.get("options", []) or []:
            add_node(option)
            for child in option.get("children", []) or []:
                add_node(child)

    for node in nodes or []:
        add_node(node)
    return index


def pop_preserved_annotations(
    node: dict[str, Any],
    annotation_index: dict[str, dict[tuple[Any, ...], list[dict[str, Any]]]],
) -> dict[str, Any] | None:
    source_key = node_fingerprint(node)
    if source_key[0]:
        source_candidates = annotation_index.get("source", {}).get(source_key) or []
        if source_candidates:
            return source_candidates.pop(0)

    content_key = node_content_fingerprint(node)
    content_candidates = annotation_index.get("content", {}).get(content_key) or []
    if content_candidates:
        return content_candidates.pop(0)
    return None


def apply_preserved_annotations(nodes: list[dict[str, Any]], annotation_index: dict[str, dict[tuple[Any, ...], list[dict[str, Any]]]]) -> None:
    def apply_to_node(node: dict[str, Any]) -> None:
        if not isinstance(node, dict):
            return
        preserved = pop_preserved_annotations(node, annotation_index)
        if preserved:
            node["annotations"] = preserved
        for option in node.get("options", []) or []:
            apply_to_node(option)
            for child in option.get("children", []) or []:
                apply_to_node(child)

    for node in nodes or []:
        apply_to_node(node)


def load_preserved_annotation_index(output_path: Path) -> dict[str, dict[tuple[Any, ...], list[dict[str, Any]]]]:
    if not output_path.exists():
        return {"source": {}, "content": {}}
    try:
        old_document = json.loads(output_path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError):
        return {"source": {}, "content": {}}
    blocks = old_document.get("blocks", []) if isinstance(old_document, dict) else []
    return build_annotation_index(blocks if isinstance(blocks, list) else [])


def normalize_header(value: Any) -> str:
    return str(value or "").strip().replace(" ", "")


def find_sheet(workbook: Any, name: str) -> Any:
    for sheet in workbook.worksheets:
        if sheet.title == name:
            return sheet
    raise ValueError(f"missing sheet: {name}")


def header_indexes(sheet: Any) -> dict[str, int]:
    headers = {normalize_header(sheet.cell(1, col).value): col for col in range(1, sheet.max_column + 1)}
    return {
        "source_id": headers.get("\u6e90ID") or headers.get("ID") or headers.get("Id") or 1,
        "order": headers.get("\u5e8f\u53f7") or headers.get("\u987a\u5e8f") or 2,
        "type": headers.get("\u5757\u7c7b\u578b") or 3,
        "speaker": headers.get("\u53d1\u8a00\u8005") or 4,
        "text": headers.get("\u6b63\u6587") or headers.get("Content") or 5,
        "keywords": headers.get("\u5173\u952e\u8bcd") or headers.get("\u6587\u672c\u5757\u5173\u952e\u8bcd") or (6 if sheet.max_column >= 6 else 0),
        "note": headers.get("\u6ce8\u91ca") or headers.get("\u5907\u6ce8") or (7 if sheet.max_column >= 7 else 0),
    }


def build_merged_cell_lookup(sheet: Any) -> dict[tuple[int, int], tuple[Any, bool]]:
    lookup: dict[tuple[int, int], tuple[Any, bool]] = {}
    for merged_range in sheet.merged_cells.ranges:
        value = sheet.cell(merged_range.min_row, merged_range.min_col).value
        for row_number in range(merged_range.min_row, merged_range.max_row + 1):
            for col_number in range(merged_range.min_col, merged_range.max_col + 1):
                lookup[(row_number, col_number)] = (value, row_number == merged_range.max_row)
    return lookup


def cell_value_with_merge_info(
    sheet: Any,
    row_number: int,
    col_number: int,
    merged_lookup: dict[tuple[int, int], tuple[Any, bool]],
) -> tuple[Any, bool]:
    if not col_number:
        return None, False
    if (row_number, col_number) in merged_lookup:
        return merged_lookup[(row_number, col_number)]
    return sheet.cell(row_number, col_number).value, True


def parse_keywords(value: Any) -> list[str]:
    seen: set[str] = set()
    keywords: list[str] = []
    for keyword in str(value or "").split("/"):
        normalized = keyword.strip()
        if normalized and normalized not in seen:
            seen.add(normalized)
            keywords.append(normalized)
    return keywords


def is_empty_row(values: list[Any]) -> bool:
    return not any(str(value or "").strip() for value in values)


def iter_rows(sheet: Any) -> list[Row]:
    indexes = header_indexes(sheet)
    merged_lookup = build_merged_cell_lookup(sheet)
    rows: list[Row] = []
    section_index = 0
    section_has_content = False
    for row_number in range(2, sheet.max_row + 1):
        source_id = str(sheet.cell(row_number, indexes["source_id"]).value or "").strip()
        raw_order = sheet.cell(row_number, indexes["order"]).value
        raw_type = str(sheet.cell(row_number, indexes["type"]).value or "").strip()
        speaker = str(sheet.cell(row_number, indexes["speaker"]).value or "").strip()
        raw_text = str(sheet.cell(row_number, indexes["text"]).value or "").strip()
        keyword_value, _ = cell_value_with_merge_info(sheet, row_number, indexes["keywords"], merged_lookup)
        note_value, is_note_anchor = cell_value_with_merge_info(sheet, row_number, indexes["note"], merged_lookup)
        keywords = parse_keywords(keyword_value)
        inline_note = str(note_value or "").strip() if is_note_anchor else ""
        if is_empty_row([source_id, raw_type, speaker, raw_text]):
            if section_has_content:
                section_index += 1
                section_has_content = False
            continue
        if not raw_type:
            raise ValueError(f"row {row_number}: missing block type")
        block_type = TYPE_MAP.get(raw_type)
        if not block_type:
            raise ValueError(f"row {row_number}: unsupported block type {raw_type!r}")
        try:
            order = int(raw_order)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"row {row_number}: invalid order {raw_order!r}") from exc
        rows.append(Row(row_number, source_id, order, raw_type, block_type, speaker, raw_text, keywords, section_index, inline_note))
        section_has_content = True
    rows.sort(key=lambda item: item.order)
    return rows


def strip_markup_to_text(value: str) -> str:
    output: list[str] = []
    cursor = 0
    for match in TE_PATTERN.finditer(value):
        output.append(TAG_PATTERN.sub("", value[cursor : match.start()]))
        term = TAG_PATTERN.sub("", match.group(3)).strip()
        output.append(term)
        cursor = match.end()
    output.append(TAG_PATTERN.sub("", value[cursor:]))
    return "".join(output)


def block_metadata(row: Row) -> dict[str, Any]:
    meta: dict[str, Any] = {
        "source_order": row.order,
        "source_type": row.raw_type,
        "source_row": row.row_number,
        "xlsx_section_index": row.section_index,
    }
    if row.source_id:
        meta["source_id"] = row.source_id
    if row.keywords:
        meta["keywords"] = row.keywords
    return meta


def make_text_block(row: Row, block_type: str | None = None) -> dict[str, Any]:
    text = strip_markup_to_text(row.raw_text)
    target_type = block_type or row.block_type
    block: dict[str, Any] = {
        "id": new_id(),
        "type": target_type,
        "text": text,
        "annotations": empty_annotations(),
        "style": {"italic": False, "speech_mode": ""},
        "meta": block_metadata(row),
    }
    if target_type == "line":
        block["speaker"] = row.speaker
    return block


def make_line_like_child(row: Row) -> dict[str, Any]:
    if row.speaker:
        return make_text_block(row, "line")
    return make_text_block(row, "narration")


def make_inline_note_block(row: Row) -> dict[str, Any]:
    meta = block_metadata(row)
    meta.pop("keywords", None)
    source_id = str(meta.get("source_id", "")).strip()
    if source_id:
        meta["inline_note_for_source_id"] = source_id
        meta["source_id"] = f"{source_id}#note"
    else:
        meta["source_id"] = f"row{row.row_number}#note"
    meta["source_type"] = "\u6ce8\u91ca"
    return {
        "id": new_id(),
        "type": "note",
        "text": strip_markup_to_text(row.inline_note),
        "annotations": empty_annotations(),
        "style": {"italic": False, "speech_mode": ""},
        "meta": meta,
    }


def append_inline_note_if_present(blocks: list[dict[str, Any]], row: Row) -> None:
    if row.inline_note:
        blocks.append(make_inline_note_block(row))


def make_branch(rows: list[Row], start_index: int) -> tuple[dict[str, Any], int]:
    first = rows[start_index]
    branch: dict[str, Any] = {
        "id": new_id(),
        "type": "branch",
        "speaker": first.speaker or "\u6f02\u6cca\u8005",
        "options": [],
        "meta": block_metadata(first),
    }
    index = start_index
    while index < len(rows) and rows[index].block_type == "branch_option":
        option_row = rows[index]
        label = strip_markup_to_text(option_row.raw_text)
        option = {
            "id": new_id("opt"),
            "label": label,
            "annotations": empty_annotations(),
            "children": [],
            "meta": block_metadata(option_row),
        }
        index += 1
        while index < len(rows) and rows[index].block_type == "branch_followup":
            option["children"].append(make_line_like_child(rows[index]))
            index += 1
        branch["options"].append(option)
    if not branch["options"]:
        branch["options"].append({"id": new_id("opt"), "label": "", "annotations": empty_annotations(), "children": []})
        index = start_index + 1
    return branch, index


def convert_rows_to_blocks(rows: list[Row]) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    index = 0
    while index < len(rows):
        row = rows[index]
        if row.block_type == "branch_option":
            branch, index = make_branch(rows, index)
            blocks.append(branch)
            append_inline_note_if_present(blocks, row)
            continue
        if row.block_type == "branch_followup":
            blocks.append(make_line_like_child(row))
        else:
            blocks.append(make_text_block(row))
        append_inline_note_if_present(blocks, row)
        index += 1
    return blocks


def relative_to_raw(path: Path, raw_root: Path) -> str:
    try:
        return path.resolve().relative_to(raw_root.resolve()).as_posix()
    except ValueError:
        return path.name


def convert_workbook(input_path: Path, output_path: Path, raw_root: Path = RAW_ROOT) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(input_path, data_only=False)
    sheet = find_sheet(workbook, ANNOTATION_SHEET_NAME)
    rows = iter_rows(sheet)
    title = output_path.name[: -len(".editor.json")] if output_path.name.endswith(".editor.json") else output_path.stem
    blocks = convert_rows_to_blocks(rows)
    preserved_annotation_index = load_preserved_annotation_index(output_path)
    apply_preserved_annotations(blocks, preserved_annotation_index)
    document = {
        "version": 5,
        "path": relative_to_raw(output_path, raw_root),
        "title": title,
        "doc_type": "story",
        "meta": {
            "source": "xlsx_annotation_table",
            "source_workbook": relative_to_raw(input_path, raw_root),
            "export_txt": False,
            "player_placeholders": {
                "PlayerName": "\u73a9\u5bb6\u7528\u6237\u540d",
                "TA": "\u73a9\u5bb6\u9009\u62e9\u7684\u6027\u522b",
            },
        },
        "blocks": blocks,
    }
    output_path.write_text(json.dumps(document, ensure_ascii=False, indent=2), encoding="utf-8")
    return document


def default_output_path(input_path: Path) -> Path:
    name = input_path.name
    for suffix in (".xlsx", ".xlsm", ".xlsl"):
        if name.lower().endswith(suffix):
            return input_path.with_name(name[: -len(suffix)] + ".editor.json")
    return input_path.with_suffix(".editor.json")


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert annotated story workbook to .editor.json")
    parser.add_argument("input", type=Path)
    parser.add_argument("-o", "--output", type=Path)
    parser.add_argument("--raw-root", type=Path, default=RAW_ROOT)
    args = parser.parse_args()
    output_path = args.output or default_output_path(args.input)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = convert_workbook(args.input, output_path, raw_root=args.raw_root)
    print(f"wrote {output_path}")
    print(f"blocks: {len(document['blocks'])}")


if __name__ == "__main__":
    main()
