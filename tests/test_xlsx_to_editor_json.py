import json

import openpyxl

from src.xlsx_to_editor_json import convert_workbook


def write_workbook(path, rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "标注表"
    sheet.append(["源ID", "序号", "块类型", "发言者", "正文", "关键词", "注释"])
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def test_convert_workbook_preserves_omissions_for_unchanged_blocks(tmp_path):
    workbook_path = tmp_path / "story.xlsx"
    output_path = tmp_path / "story.editor.json"
    write_workbook(
        workbook_path,
        [
            ["A1", 1, "旁白", "", "来自罗伊族的罗伊族人。"],
            ["A2", 2, "台词", "漂泊者", "这里是新文本。"],
        ],
    )

    convert_workbook(workbook_path, output_path, raw_root=tmp_path)
    document = json.loads(output_path.read_text(encoding="utf-8"))
    document["blocks"][0]["annotations"]["omissions"] = [{"term": "罗伊族", "start": 6, "end": 9}]
    output_path.write_text(json.dumps(document, ensure_ascii=False, indent=2), encoding="utf-8")

    write_workbook(
        workbook_path,
        [
            ["A2", 1, "台词", "漂泊者", "这里是新文本。"],
            ["A1", 2, "旁白", "", "来自罗伊族的罗伊族人。"],
        ],
    )
    converted = convert_workbook(workbook_path, output_path, raw_root=tmp_path)

    narration = next(block for block in converted["blocks"] if block["meta"]["source_id"] == "A1")
    assert narration["annotations"]["omissions"] == [{"term": "罗伊族", "start": 6, "end": 9}]


def test_convert_workbook_drops_omissions_when_block_text_changes(tmp_path):
    workbook_path = tmp_path / "story.xlsx"
    output_path = tmp_path / "story.editor.json"
    write_workbook(workbook_path, [["A1", 1, "旁白", "", "来自罗伊族的罗伊族人。"]])

    convert_workbook(workbook_path, output_path, raw_root=tmp_path)
    document = json.loads(output_path.read_text(encoding="utf-8"))
    document["blocks"][0]["annotations"]["omissions"] = [{"term": "罗伊族", "start": 6, "end": 9}]
    output_path.write_text(json.dumps(document, ensure_ascii=False, indent=2), encoding="utf-8")

    write_workbook(workbook_path, [["A1", 1, "旁白", "", "来自罗伊族的战士。"]])
    converted = convert_workbook(workbook_path, output_path, raw_root=tmp_path)

    assert converted["blocks"][0]["annotations"]["omissions"] == []


def test_convert_workbook_preserves_omissions_from_legacy_blocks_without_source_id(tmp_path):
    workbook_path = tmp_path / "story.xlsx"
    output_path = tmp_path / "story.editor.json"
    text = "游戏中那个举剑扇动羽翼的主角形象，让你感到了几分熟悉。"
    write_workbook(workbook_path, [["MAIN_YHX_12162_4", 1, "旁白", "", text]])

    legacy_document = {
        "version": 4,
        "blocks": [
            {
                "id": "blk_old",
                "type": "narration",
                "text": text,
                "annotations": {"manual_refs": [], "omissions": [{"term": "角", "start": 13, "end": 14}], "settings": []},
            }
        ],
    }
    output_path.write_text(json.dumps(legacy_document, ensure_ascii=False, indent=2), encoding="utf-8")

    converted = convert_workbook(workbook_path, output_path, raw_root=tmp_path)

    assert converted["blocks"][0]["annotations"]["omissions"] == [{"term": "角", "start": 13, "end": 14}]


def test_convert_workbook_reads_keywords_and_merged_inline_notes(tmp_path):
    workbook_path = tmp_path / "story.xlsx"
    output_path = tmp_path / "story.editor.json"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "标注表"
    sheet.append(["源ID", "序号", "块类型", "发言者", "正文", "关键词", "注释"])
    sheet.append(["A312", 312, "台词", "爱弥斯", "第一句"])
    sheet.append(["A313", 313, "台词", "爱弥斯", "第二句"])
    sheet.append(["A314", 314, "台词", "爱弥斯", "第三句"])
    sheet.append(["A321", 321, "旁白", "", "第四句"])
    sheet.append(["A322", 322, "旁白", "", "第五句"])
    sheet.append(["A323", 323, "旁白", "", "第六句"])
    sheet.append(["A324", 324, "旁白", "", "第七句"])
    sheet.merge_cells(start_row=2, start_column=6, end_row=4, end_column=6)
    sheet.cell(2, 6).value = "陆·赫斯/星炬学院印象"
    sheet.merge_cells(start_row=5, start_column=7, end_row=8, end_column=7)
    sheet.cell(5, 7).value = "课堂上漂泊者和爱弥斯说话，其他人视角他在自言自语"
    workbook.save(workbook_path)

    converted = convert_workbook(workbook_path, output_path, raw_root=tmp_path)

    keyed_blocks = [block for block in converted["blocks"] if block["meta"].get("source_id") in {"A312", "A313", "A314"}]
    assert [block["meta"]["keywords"] for block in keyed_blocks] == [
        ["陆·赫斯", "星炬学院印象"],
        ["陆·赫斯", "星炬学院印象"],
        ["陆·赫斯", "星炬学院印象"],
    ]
    source_ids = [block["meta"].get("source_id") for block in converted["blocks"]]
    assert source_ids[source_ids.index("A324") + 1] == "A324#note"
    assert converted["blocks"][source_ids.index("A324#note")]["text"] == "课堂上漂泊者和爱弥斯说话，其他人视角他在自言自语"


def test_convert_workbook_keeps_blank_row_sections_and_text_blocks_without_source_ids(tmp_path):
    workbook_path = tmp_path / "role_story.xlsx"
    output_path = tmp_path / "role_story.editor.json"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "标注表"
    sheet.append(["源ID", "序号", "块类型", "发言者", "正文", "关键词", "注释"])
    sheet.append([None, 1, "文本标题", None, "标题一"])
    sheet.append([None, 2, "文本正文", None, "正文一"])
    sheet.append([None, 3, None, None, None])
    sheet.append([None, 4, "文本标题", None, "标题二"])
    sheet.append([None, 5, "文本正文", None, "正文二"])
    sheet.merge_cells(start_row=2, start_column=6, end_row=4, end_column=6)
    sheet.cell(2, 6).value = "第一组关键词"
    workbook.save(workbook_path)

    converted = convert_workbook(workbook_path, output_path, raw_root=tmp_path)

    assert [block["type"] for block in converted["blocks"]] == ["text_title", "text_body", "text_title", "text_body"]
    assert [block["meta"]["xlsx_section_index"] for block in converted["blocks"]] == [0, 0, 1, 1]
    assert [block["meta"]["source_order"] for block in converted["blocks"]] == [1, 2, 4, 5]
    assert [block["meta"].get("keywords") for block in converted["blocks"][:2]] == [["第一组关键词"], ["第一组关键词"]]


def test_convert_workbook_rejects_content_rows_without_order(tmp_path):
    workbook_path = tmp_path / "role_story.xlsx"
    output_path = tmp_path / "role_story.editor.json"
    write_workbook(workbook_path, [[None, None, "文本正文", None, "正文"]])

    try:
        convert_workbook(workbook_path, output_path, raw_root=tmp_path)
    except ValueError as exc:
        assert "invalid order" in str(exc)
    else:
        raise AssertionError("missing order should be rejected")
